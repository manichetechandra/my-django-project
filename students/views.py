from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import login, authenticate, logout, update_session_auth_hash
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.db import transaction
from django.views.decorators.cache import never_cache
from django.views.decorators.http import require_http_methods
from django.http import JsonResponse, HttpResponse
from django.http import JsonResponse
from django.db.models import Q, Count  # ← جديد: للبحث المتقدم والإحصائيات
from django.utils import timezone  # ← جديد: للتاريخ والوقت
from datetime import timedelta  # ← جديد: لحساب الفترات الزمنية


# ========================================
# جديد: imports للـ Excel
# ========================================
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


from .decorators import staff_required  # ← جديد: الديكوريتور الذي أنشأناه
from .forms import (
    StudentRegistrationForm, 
    StudentLoginForm,
    ApplicationSubmissionForm,
    ApplicationDocumentForm,
    StudentProfileUpdateForm,  # ← جديد
    UserEmailUpdateForm,  # ← جديد
    PasswordChangeForm  # ← جديد
    
)
from .models import (
    StudentProfile, 
    StudentApplication, 
    ApplicationType, 
    StudentNotification,
    ApplicationDocument
)


def student_register(request):
    """صفحة تسجيل طالب جديد"""
    
    if request.user.is_authenticated:
        return redirect('students:dashboard')
    
    if request.method == 'POST':
        form = StudentRegistrationForm(request.POST)
        
        if form.is_valid():
            try:
                with transaction.atomic():
                    # إنشاء User
                    user = form.save()
                    
                    # إنشاء Student Profile
                    student_profile = StudentProfile.objects.create(
                        user=user,
                        registration_number=form.cleaned_data['registration_number'],
                        full_name_arabic=form.cleaned_data['full_name_arabic'],
                        birth_date=form.cleaned_data['birth_date'],
                        birth_place=form.cleaned_data['birth_place'],
                        gender=form.cleaned_data['gender'],
                        phone=form.cleaned_data['phone'],
                        address=form.cleaned_data['address'],
                        wilaya=form.cleaned_data['wilaya'],
                        faculty=form.cleaned_data['faculty'],
                        department=form.cleaned_data['department'],
                        current_year=form.cleaned_data['current_year'],
                        is_verified=True,
                    )
                    
                    # إنشاء إشعار ترحيبي
                    StudentNotification.objects.create(
                        student=student_profile,
                        title='مرحباً بك!',
                        message=f'مرحباً {student_profile.full_name_arabic}، تم إنشاء حسابك بنجاح. يمكنك الآن تقديم الطلبات.',
                        link='/students/dashboard/'
                    )
                    
                    # تسجيل الدخول تلقائياً
                    login(request, user)
                    
                    messages.success(request, 'تم إنشاء حسابك بنجاح! مرحباً بك')
                    return redirect('students:dashboard')
                    
            except Exception as e:
                messages.error(request, f'حدث خطأ: {str(e)}')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f'{error}')
    else:
        form = StudentRegistrationForm()
    
    context = {
        'form': form,
        'page_title': 'تسجيل طالب جديد'
    }
    return render(request, 'students/register.html', context)


@never_cache
def student_login(request):
    """صفحة تسجيل الدخول"""
    
    if request.user.is_authenticated:
        # توجيه ذكي: إذا مشرف → لوحة المشرف، إذا طالب → لوحة الطالب
        if request.user.is_staff:
            return redirect('students:admin_dashboard')
        else:
            return redirect('students:dashboard')
    
    if request.method == 'POST':
        form = StudentLoginForm(request.POST)
        
        if form.is_valid():
            username = form.cleaned_data['username']
            password = form.cleaned_data['password']
            remember_me = form.cleaned_data.get('remember_me', False)
            
            user = authenticate(request, username=username, password=password)
            
            if user is not None:
                # تسجيل الدخول
                login(request, user)
                
                if not remember_me:
                    request.session.set_expiry(0)
                
                # توجيه ذكي حسب نوع المستخدم
                if user.is_staff:
                    messages.success(request, f'مرحباً! تم تسجيل الدخول كمشرف')
                    next_url = request.GET.get('next', 'students:admin_dashboard')
                else:
                    try:
                        student_profile = user.student_profile
                        
                        if not student_profile.is_active:
                            messages.error(request, 'حسابك غير نشط. يرجى الاتصال بالإدارة')
                            logout(request)
                            return redirect('students:login')
                        
                        messages.success(request, f'مرحباً {student_profile.full_name_arabic}!')
                        next_url = request.GET.get('next', 'students:dashboard')
                    except StudentProfile.DoesNotExist:
                        messages.warning(request, 'لا يوجد ملف طالب. يمكنك الوصول للوحة الإدارة فقط')
                        next_url = 'students:admin_dashboard' if user.is_staff else 'homepage:index'
                
                response = redirect(next_url)
                response['Cache-Control'] = 'no-cache, no-store, must-revalidate, private'
                return response
            else:
                messages.error(request, 'اسم المستخدم أو كلمة المرور غير صحيحة')
    else:
        form = StudentLoginForm()
    
    context = {
        'form': form,
        'page_title': 'تسجيل الدخول'
    }
    
    response = render(request, 'students/login.html', context)
    response['Cache-Control'] = 'no-cache, no-store, must-revalidate, private'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'
    
    return response
@never_cache
def student_logout(request):
    """تسجيل الخروج"""
    logout(request)
    messages.success(request, 'تم تسجيل الخروج بنجاح')
    
    response = redirect('students:login')
    response.delete_cookie('sessionid')
    response['Cache-Control'] = 'no-cache, no-store, must-revalidate, private'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'
    
    return response


@never_cache
@login_required(login_url='students:login')
def student_dashboard(request):
    """لوحة تحكم الطالب"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # إحصائيات
    applications = student.applications.all()
    pending_count = applications.filter(status='pending').count()
    under_review_count = applications.filter(status='under_review').count()
    approved_count = applications.filter(status='approved').count()
    rejected_count = applications.filter(status='rejected').count()
    
    # آخر الطلبات
    recent_applications = applications.order_by('-submitted_at')[:5]
    
    # الإشعارات غير المقروءة
    unread_notifications = student.student_notifications.filter(is_read=False)[:5]
    
    # أنواع الطلبات المتاحة
    available_application_types = ApplicationType.objects.filter(is_active=True)
    
    context = {
        'student': student,
        'pending_count': pending_count,
        'under_review_count': under_review_count,
        'approved_count': approved_count,
        'rejected_count': rejected_count,
        'recent_applications': recent_applications,
        'unread_notifications': unread_notifications,
        'available_application_types': available_application_types,
        'page_title': 'لوحة التحكم'
    }
    
    response = render(request, 'students/dashboard.html', context)
    response['Cache-Control'] = 'no-cache, no-store, must-revalidate, private'
    response['Pragma'] = 'no-cache'
    response['Expires'] = '0'
    
    return response


@never_cache
@login_required(login_url='students:login')
def create_application(request, app_type_id=None):
    """صفحة تقديم طلب جديد"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # إذا تم اختيار نوع الطلب مسبقاً
    selected_app_type = None
    initial_data = {}
    if app_type_id:
        selected_app_type = get_object_or_404(ApplicationType, id=app_type_id, is_active=True)
        initial_data['application_type'] = selected_app_type
    
    if request.method == 'POST':
        form = ApplicationSubmissionForm(request.POST)
        
        if form.is_valid():
            try:
                with transaction.atomic():
                    # إنشاء الطلب
                    application = form.save(commit=False)
                    application.student = student
                    application.status = 'pending'
                    application.save()
                    
                    # معالجة الوثائق المرفقة
                    files = request.FILES.getlist('documents')
                    file_names = request.POST.getlist('document_names')
                    
                    for i, file in enumerate(files):
                        if file:
                            doc_name = file_names[i] if i < len(file_names) else file.name
                            ApplicationDocument.objects.create(
                                application=application,
                                document_name=doc_name,
                                document_file=file
                            )
                    
                    # إنشاء إشعار
                    StudentNotification.objects.create(
                        student=student,
                        title='تم تقديم طلبك بنجاح',
                        message=f'تم تقديم طلب {application.application_type.name} برقم {application.application_number}',
                        link=f'/students/applications/{application.id}/'
                    )
                    
                    messages.success(request, f'تم تقديم طلبك بنجاح! رقم الطلب: {application.application_number}')
                    return redirect('students:application_detail', pk=application.id)
                    
            except Exception as e:
                messages.error(request, f'حدث خطأ: {str(e)}')
        else:
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f'{error}')
    else:
        form = ApplicationSubmissionForm(initial=initial_data)
    
    # جلب كل أنواع الطلبات النشطة
    application_types = ApplicationType.objects.filter(is_active=True)
    
    context = {
        'form': form,
        'application_types': application_types,
        'page_title': 'تقديم طلب جديد',
        'selected_type': selected_app_type
    }
    
    return render(request, 'students/create_application.html', context)


@never_cache
@login_required(login_url='students:login')
def my_applications(request):
    """صفحة طلباتي - قائمة كل طلبات الطالب"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # جلب كل الطلبات
    applications = student.applications.all().order_by('-submitted_at')
    
    # الفلترة حسب الحالة
    status_filter = request.GET.get('status', '')
    if status_filter:
        applications = applications.filter(status=status_filter)
    
    # الفلترة حسب النوع
    type_filter = request.GET.get('type', '')
    if type_filter:
        applications = applications.filter(application_type_id=type_filter)
    
    # إحصائيات
    total_count = student.applications.count()
    pending_count = student.applications.filter(status='pending').count()
    under_review_count = student.applications.filter(status='under_review').count()
    approved_count = student.applications.filter(status='approved').count()
    rejected_count = student.applications.filter(status='rejected').count()
    
    # أنواع الطلبات للفلتر
    application_types = ApplicationType.objects.filter(is_active=True)
    
    context = {
        'student': student,
        'applications': applications,
        'total_count': total_count,
        'pending_count': pending_count,
        'under_review_count': under_review_count,
        'approved_count': approved_count,
        'rejected_count': rejected_count,
        'application_types': application_types,
        'status_filter': status_filter,
        'type_filter': type_filter,
        'page_title': 'طلباتي'
    }
    
    return render(request, 'students/my_applications.html', context)


@never_cache
@login_required(login_url='students:login')
def application_detail(request, pk):
    """صفحة تفاصيل الطلب"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # جلب الطلب وتأكد أنه يخص الطالب
    application = get_object_or_404(
        StudentApplication, 
        id=pk, 
        student=student
    )
    
    # جلب الوثائق المرفقة
    documents = application.documents.all()
    
    context = {
        'student': student,
        'application': application,
        'documents': documents,
        'page_title': f'تفاصيل الطلب - {application.application_number}'
    }
    
    return render(request, 'students/application_detail.html', context)


@login_required(login_url='students:login')
@require_http_methods(["POST"])
def cancel_application(request, pk):
    """إلغاء طلب"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    application = get_object_or_404(
        StudentApplication, 
        id=pk, 
        student=student
    )
    
    # يمكن إلغاء الطلب فقط إذا كان في حالة "قيد الانتظار"
    if application.status == 'pending':
        application.status = 'cancelled'
        application.save()
        
        # إنشاء إشعار
        StudentNotification.objects.create(
            student=student,
            title='تم إلغاء الطلب',
            message=f'تم إلغاء طلب {application.application_type.name} برقم {application.application_number}',
        )
        
        messages.success(request, 'تم إلغاء الطلب بنجاح')
    else:
        messages.error(request, 'لا يمكن إلغاء هذا الطلب')
    
    return redirect('students:my_applications')


def csrf_failure(request, reason=""):
    """معالجة أخطاء CSRF"""
    return render(request, '403_csrf.html', status=403)


@never_cache
@login_required(login_url='students:login')
def application_print(request, pk):
    """نسخة للطباعة من تفاصيل الطلب"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    application = get_object_or_404(
        StudentApplication, 
        id=pk, 
        student=student
    )
    
    documents = application.documents.all()
    
    context = {
        'student': student,
        'application': application,
        'documents': documents,
    }
    
    return render(request, 'students/application_print.html', context)





# ========================================
# Views خاصة بالمشرف/الإدارة
# ========================================

@never_cache
@staff_required
def admin_dashboard(request):
    """
    لوحة تحكم المشرف - الصفحة الرئيسية للإدارة
    """
    
    # ==========================================
    # 1. إحصائيات عامة - كل الطلبات
    # ==========================================
    
    total_applications = StudentApplication.objects.count()
    
    pending_count = StudentApplication.objects.filter(status='pending').count()
    under_review_count = StudentApplication.objects.filter(status='under_review').count()
    approved_count = StudentApplication.objects.filter(status='approved').count()
    rejected_count = StudentApplication.objects.filter(status='rejected').count()
    
    # ==========================================
    # 2. إحصائيات زمنية - اليوم/الأسبوع/الشهر
    # ==========================================
    
    now = timezone.now()
    today_start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    week_start = now - timedelta(days=7)
    month_start = now - timedelta(days=30)
    
    today_applications = StudentApplication.objects.filter(
        submitted_at__gte=today_start
    ).count()
    
    week_applications = StudentApplication.objects.filter(
        submitted_at__gte=week_start
    ).count()
    
    month_applications = StudentApplication.objects.filter(
        submitted_at__gte=month_start
    ).count()
    
    # ==========================================
    # 3. آخر الطلبات المقدمة (آخر 10)
    # ==========================================
    
    recent_applications = StudentApplication.objects.select_related(
        'student',
        'application_type',
        'reviewed_by'
    ).order_by('-submitted_at')[:10]
    
    # ==========================================
    # 4. إحصائيات حسب نوع الطلب
    # ==========================================
    
    applications_by_type = ApplicationType.objects.annotate(
        total=Count('studentapplication')
    ).filter(total__gt=0)
    
    # ==========================================
    # 5. عدد الطلبة المسجلين
    # ==========================================
    
    total_students = StudentProfile.objects.filter(is_active=True).count()
    
    # ==========================================
    # 6. بيانات الرسوم البيانية ← جديد! 📊
    # ==========================================
    
    # ----------------------------------------
    # أ) بيانات Pie Chart - حسب الحالة
    # ----------------------------------------
    
    chart_status_labels = ['قيد الانتظار', 'قيد المراجعة', 'معتمد', 'مرفوض']
    chart_status_data = [pending_count, under_review_count, approved_count, rejected_count]
    chart_status_colors = ['#f39c12', '#3498db', '#2ecc71', '#e74c3c']
    
    # ----------------------------------------
    # ب) بيانات Bar Chart - حسب النوع
    # ----------------------------------------
    
    chart_type_labels = []
    chart_type_data = []
    chart_type_colors = []
    
    for app_type in applications_by_type:
        chart_type_labels.append(app_type.name)
        chart_type_data.append(app_type.total)
        chart_type_colors.append(app_type.color or '#3498db')
    
    # ----------------------------------------
    # ج) بيانات Line Chart - آخر 7 أيام
    # ----------------------------------------
    
    chart_days_labels = []
    chart_days_data = []
    
    for i in range(6, -1, -1):  # من 6 إلى 0 (آخر 7 أيام)
        day = now - timedelta(days=i)
        day_start = day.replace(hour=0, minute=0, second=0, microsecond=0)
        day_end = day_start + timedelta(days=1)
        
        # عدد الطلبات في هذا اليوم
        day_count = StudentApplication.objects.filter(
            submitted_at__gte=day_start,
            submitted_at__lt=day_end
        ).count()
        
        # اسم اليوم (مثلاً: "الأحد 19/11")
        day_name = day.strftime('%a %d/%m')  # Mon 19/11
        
        chart_days_labels.append(day_name)
        chart_days_data.append(day_count)
    
    # ==========================================
    # 7. تجهيز الـ Context
    # ==========================================
    
    context = {
        # الإحصائيات العامة
        'total_applications': total_applications,
        'pending_count': pending_count,
        'under_review_count': under_review_count,
        'approved_count': approved_count,
        'rejected_count': rejected_count,
        
        # الإحصائيات الزمنية
        'today_applications': today_applications,
        'week_applications': week_applications,
        'month_applications': month_applications,
        
        # القوائم
        'recent_applications': recent_applications,
        'applications_by_type': applications_by_type,
        
        # أخرى
        'total_students': total_students,
        'page_title': 'لوحة تحكم المشرف',
        
        # ==========================================
        # بيانات الرسوم البيانية ← جديد!
        # ==========================================
        'chart_status_labels': chart_status_labels,
        'chart_status_data': chart_status_data,
        'chart_status_colors': chart_status_colors,
        
        'chart_type_labels': chart_type_labels,
        'chart_type_data': chart_type_data,
        'chart_type_colors': chart_type_colors,
        
        'chart_days_labels': chart_days_labels,
        'chart_days_data': chart_days_data,
    }
    
    return render(request, 'students/admin/dashboard.html', context)


@never_cache
@staff_required
def admin_applications_list(request):
    """
    صفحة قائمة كل الطلبات للمشرف
    
    ماذا تفعل؟
    1. تجلب كل الطلبات
    2. تطبق الفلترة (حسب الحالة، النوع، البحث)
    3. تعرضها في جدول منظم
    4. تدعم الإجراءات الجماعية (Bulk Actions) ← جديد!
    """
    
    # ==========================================
    # 1. جلب كل الطلبات مع معلومات مرتبطة
    # ==========================================
    
    applications = StudentApplication.objects.select_related(
        'student__user',
        'application_type',
        'reviewed_by'
    ).order_by('-submitted_at')
    
    # ==========================================
    # 2. الفلترة حسب الحالة
    # ==========================================
    
    status_filter = request.GET.get('status', '')
    
    if status_filter:
        applications = applications.filter(status=status_filter)
    
    # ==========================================
    # 3. الفلترة حسب نوع الطلب
    # ==========================================
    
    type_filter = request.GET.get('type', '')
    
    if type_filter:
        applications = applications.filter(application_type_id=type_filter)
    
    # ==========================================
    # 4. البحث (في رقم الطلب أو اسم الطالب)
    # ==========================================
    
    search_query = request.GET.get('search', '')
    
    if search_query:
        applications = applications.filter(
            Q(application_number__icontains=search_query) |
            Q(student__full_name_arabic__icontains=search_query) |
            Q(student__registration_number__icontains=search_query)
        )
    
    # ==========================================
    # 5. الإجراءات الجماعية (Bulk Actions) ← جديد!
    # ==========================================
    
    if request.method == 'POST':
        # جلب الإجراء المطلوب
        bulk_action = request.POST.get('bulk_action')
        # جلب IDs الطلبات المحددة
        selected_ids = request.POST.getlist('selected_applications')
        
        if not selected_ids:
            messages.warning(request, 'لم تقم بتحديد أي طلبات!')
        elif not bulk_action:
            messages.warning(request, 'لم تقم باختيار إجراء!')
        else:
            try:
                with transaction.atomic():
                    # جلب الطلبات المحددة
                    selected_applications = StudentApplication.objects.filter(
                        id__in=selected_ids
                    ).select_related('student', 'application_type')
                    
                    count = selected_applications.count()
                    
                    # ==========================================
                    # الإجراء: اعتماد الكل
                    # ==========================================
                    if bulk_action == 'approve':
                        for app in selected_applications:
                            app.status = 'approved'
                            app.reviewed_by = request.user
                            app.reviewed_at = timezone.now()
                            app.save()
                            
                            # إنشاء إشعار للطالب
                            StudentNotification.objects.create(
                                student=app.student,
                                title='تم اعتماد طلبك ✅',
                                message=f'تم اعتماد طلب {app.application_type.name} برقم {app.application_number}',
                                link=f'/students/applications/{app.id}/'
                            )
                        
                        messages.success(request, f'✅ تم اعتماد {count} طلب بنجاح!')
                    
                    # ==========================================
                    # الإجراء: رفض الكل
                    # ==========================================
                    elif bulk_action == 'reject':
                        admin_notes = request.POST.get('bulk_notes', 'تم رفض الطلب')
                        
                        for app in selected_applications:
                            app.status = 'rejected'
                            app.reviewed_by = request.user
                            app.reviewed_at = timezone.now()
                            app.admin_notes = admin_notes
                            app.save()
                            
                            # إنشاء إشعار للطالب
                            StudentNotification.objects.create(
                                student=app.student,
                                title='تم رفض طلبك ❌',
                                message=f'تم رفض طلب {app.application_type.name} برقم {app.application_number}. يرجى مراجعة الملاحظات.',
                                link=f'/students/applications/{app.id}/'
                            )
                        
                        messages.warning(request, f'⚠️ تم رفض {count} طلب')
                    
                    # ==========================================
                    # الإجراء: قيد المراجعة
                    # ==========================================
                    elif bulk_action == 'review':
                        for app in selected_applications:
                            app.status = 'under_review'
                            app.save()
                        
                        messages.info(request, f'📝 تم تغيير حالة {count} طلب إلى "قيد المراجعة"')
                    
                    # ==========================================
                    # الإجراء: قيد الانتظار
                    # ==========================================
                    elif bulk_action == 'pending':
                        for app in selected_applications:
                            app.status = 'pending'
                            app.save()
                        
                        messages.info(request, f'⏳ تم تغيير حالة {count} طلب إلى "قيد الانتظار"')
                    
                    # ==========================================
                    # الإجراء: حذف
                    # ==========================================
                    elif bulk_action == 'delete':
                        selected_applications.delete()
                        messages.success(request, f'🗑️ تم حذف {count} طلب')
                    
                    # إعادة التوجيه لتجنب إعادة إرسال POST
                    return redirect(request.path + '?' + request.GET.urlencode())
                    
            except Exception as e:
                messages.error(request, f'حدث خطأ: {str(e)}')
    
    # ==========================================
    # 6. إحصائيات للفلتر
    # ==========================================
    
    total_count = applications.count()
    
    # ==========================================
    # 7. جلب أنواع الطلبات للفلتر
    # ==========================================
    
    application_types = ApplicationType.objects.filter(is_active=True)
    
    # ==========================================
    # 8. تجهيز الـ Context
    # ==========================================
    
    context = {
        'applications': applications,
        'total_count': total_count,
        'application_types': application_types,
        'status_filter': status_filter,
        'type_filter': type_filter,
        'search_query': search_query,
        'page_title': 'إدارة الطلبات'
    }
    
    return render(request, 'students/admin/applications_list.html', context)

@never_cache
@staff_required
def admin_application_detail(request, pk):
    """صفحة تفاصيل الطلب للمشرف"""
    
    application = get_object_or_404(
        StudentApplication.objects.select_related(
            'student__user',
            'application_type',
            'reviewed_by'
        ),
        id=pk
    )
    
    documents = application.documents.all()
    
    if request.method == 'POST':
        action = request.POST.get('action')
        admin_notes = request.POST.get('admin_notes', '')
        
        if not action:
            messages.warning(request, 'لم يتم تحديد الإجراء المطلوب')
            return redirect('students:admin_application_detail', pk=pk)
        
        try:
            with transaction.atomic():
                
                # ==========================================
                # حالة الاعتماد
                # ==========================================
                if action == 'approve':
                    application.status = 'approved'
                    application.reviewed_by = request.user
                    application.reviewed_at = timezone.now()
                    if admin_notes:
                        application.admin_notes = admin_notes
                    application.save()
                    
                    # ← إنشاء إشعار
                    StudentNotification.objects.create(
                        student=application.student,
                        title='تم اعتماد طلبك ✅',
                        message=f'تم اعتماد طلب {application.application_type.name} برقم {application.application_number}',
                        link=f'/students/applications/{application.id}/',
                        notification_type='success'
                    )
                    
                    print(f"✅ تم إنشاء إشعار اعتماد للطالب {application.student.full_name_arabic}")
                    
                    messages.success(request, 'تم اعتماد الطلب بنجاح وإرسال إشعار للطالب')
                    return redirect('students:admin_application_detail', pk=pk)
                
                # ==========================================
                # حالة الرفض
                # ==========================================
                elif action == 'reject':
                    application.status = 'rejected'
                    application.reviewed_by = request.user
                    application.reviewed_at = timezone.now()
                    application.admin_notes = admin_notes
                    application.save()
                    
                    # ← إنشاء إشعار
                    StudentNotification.objects.create(
                        student=application.student,
                        title='تم رفض طلبك ❌',
                        message=f'تم رفض طلب {application.application_type.name} برقم {application.application_number}. يرجى مراجعة الملاحظات.',
                        link=f'/students/applications/{application.id}/',
                        notification_type='error'
                    )
                    
                    print(f"✅ تم إنشاء إشعار رفض للطالب {application.student.full_name_arabic}")
                    
                    messages.warning(request, 'تم رفض الطلب وإرسال إشعار للطالب')
                    return redirect('students:admin_application_detail', pk=pk)
                
                # ==========================================
                # حالة قيد الانتظار
                # ==========================================
                elif action == 'pending':
                    application.status = 'pending'
                    if admin_notes:
                        application.admin_notes = admin_notes
                    application.save()
                    
                    # ← إنشاء إشعار
                    StudentNotification.objects.create(
                        student=application.student,
                        title='تم تغيير حالة طلبك ⏳',
                        message=f'تم تغيير حالة طلب {application.application_type.name} برقم {application.application_number} إلى "قيد الانتظار".',
                        link=f'/students/applications/{application.id}/',
                        notification_type='warning'
                    )
                    
                    print(f"✅ تم إنشاء إشعار انتظار للطالب {application.student.full_name_arabic}")
                    
                    messages.info(request, 'تم تغيير حالة الطلب إلى "قيد الانتظار" وإرسال إشعار للطالب')
                    return redirect('students:admin_application_detail', pk=pk)
                
                # ==========================================
                # حالة تحت المراجعة
                # ==========================================
                elif action == 'review':
                    application.status = 'under_review'
                    if admin_notes:
                        application.admin_notes = admin_notes
                    application.save()
                    
                    # ← إنشاء إشعار
                    StudentNotification.objects.create(
                        student=application.student,
                        title='طلبك قيد المراجعة 📝',
                        message=f'طلب {application.application_type.name} برقم {application.application_number} قيد المراجعة من قبل الإدارة.',
                        link=f'/students/applications/{application.id}/',
                        notification_type='info'
                    )
                    
                    print(f"✅ تم إنشاء إشعار مراجعة للطالب {application.student.full_name_arabic}")
                    
                    messages.info(request, 'تم تغيير حالة الطلب إلى "قيد المراجعة" وإرسال إشعار للطالب')
                    return redirect('students:admin_application_detail', pk=pk)
                
                else:
                    messages.warning(request, f'إجراء غير معروف: {action}')
                    return redirect('students:admin_application_detail', pk=pk)
                    
        except Exception as e:
            print(f"❌ خطأ في تغيير الحالة: {str(e)}")
            messages.error(request, f'حدث خطأ: {str(e)}')
            return redirect('students:admin_application_detail', pk=pk)
    
    context = {
        'application': application,
        'student': application.student,
        'documents': documents,
        'page_title': f'تفاصيل الطلب - {application.application_number}'
    }
    
    return render(request, 'students/admin/application_detail.html', context)

    # ========================================
# تصدير الطلبات إلى Excel
# ========================================

@never_cache
@staff_required
def export_applications_excel(request):
    """
    تصدير قائمة الطلبات إلى ملف Excel
    
    ماذا تفعل؟
    1. تجلب الطلبات (مع الفلترة إذا وجدت)
    2. تنشئ ملف Excel
    3. تضيف البيانات مع تنسيق جميل
    4. تُرجع الملف للتحميل
    """
    
    # ==========================================
    # 1. جلب الطلبات (نفس منطق admin_applications_list)
    # ==========================================
    
    applications = StudentApplication.objects.select_related(
        'student__user',
        'application_type',
        'reviewed_by'
    ).order_by('-submitted_at')
    
    # الفلترة حسب URL parameters
    status_filter = request.GET.get('status', '')
    if status_filter:
        applications = applications.filter(status=status_filter)
    
    type_filter = request.GET.get('type', '')
    if type_filter:
        applications = applications.filter(application_type_id=type_filter)
    
    search_query = request.GET.get('search', '')
    if search_query:
        applications = applications.filter(
            Q(application_number__icontains=search_query) |
            Q(student__full_name_arabic__icontains=search_query) |
            Q(student__registration_number__icontains=search_query)
        )
    
    # ==========================================
    # 2. إنشاء Workbook جديد
    # ==========================================
    
    wb = Workbook()
    ws = wb.active
    ws.title = "الطلبات"
    
    # ==========================================
    # 3. تنسيق العناوين
    # ==========================================
    
    # ألوان
    header_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    header_font = Font(name='Arial', size=12, bold=True, color="FFFFFF")
    
    # محاذاة
    center_alignment = Alignment(horizontal="center", vertical="center")
    right_alignment = Alignment(horizontal="right", vertical="center")
    
    # حدود
    thin_border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # ==========================================
    # 4. إضافة العناوين
    # ==========================================
    
    headers = [
        '#',
        'رقم الطلب',
        'اسم الطالب',
        'رقم التسجيل',
        'نوع الطلب',
        'الحالة',
        'تاريخ التقديم',
        'تاريخ المراجعة',
        'راجعه',
        'الهاتف',
        'البريد الإلكتروني',
        'الكلية',
        'القسم',
        'السنة الدراسية'
    ]
    
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.value = header
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_alignment
        cell.border = thin_border
    
    # ==========================================
    # 5. إضافة البيانات
    # ==========================================
    
    # ألوان الصفوف (متبادلة)
    row_fill_1 = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    row_fill_2 = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
    
    # ألوان الحالات
    status_colors = {
        'pending': PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid"),
        'under_review': PatternFill(start_color="D1ECF1", end_color="D1ECF1", fill_type="solid"),
        'approved': PatternFill(start_color="D4EDDA", end_color="D4EDDA", fill_type="solid"),
        'rejected': PatternFill(start_color="F8D7DA", end_color="F8D7DA", fill_type="solid"),
        'cancelled': PatternFill(start_color="E2E3E5", end_color="E2E3E5", fill_type="solid"),
    }
    
    for row_num, app in enumerate(applications, 2):
        # تحديد لون الصف
        row_fill = row_fill_1 if row_num % 2 == 0 else row_fill_2
        
        # البيانات
        data = [
            row_num - 1,  # الرقم التسلسلي
            app.application_number,
            app.student.full_name_arabic,
            app.student.registration_number,
            app.application_type.name,
            app.get_status_display(),
            app.submitted_at.strftime('%Y-%m-%d %H:%M') if app.submitted_at else '',
            app.reviewed_at.strftime('%Y-%m-%d %H:%M') if app.reviewed_at else '',
            app.reviewed_by.get_full_name() if app.reviewed_by else '',
            app.student.phone,
            app.student.user.email,
            app.student.faculty,
            app.student.department,
            app.student.get_current_year_display(),
        ]
        
        for col_num, value in enumerate(data, 1):
            cell = ws.cell(row=row_num, column=col_num)
            cell.value = value
            cell.border = thin_border
            cell.alignment = right_alignment if col_num > 1 else center_alignment
            
            # لون خاص لعمود الحالة
            if col_num == 6:  # عمود الحالة
                cell.fill = status_colors.get(app.status, row_fill)
            else:
                cell.fill = row_fill
    
    # ==========================================
    # 6. ضبط عرض الأعمدة
    # ==========================================
    
    column_widths = {
        'A': 8,   # #
        'B': 15,  # رقم الطلب
        'C': 25,  # اسم الطالب
        'D': 15,  # رقم التسجيل
        'E': 20,  # نوع الطلب
        'F': 15,  # الحالة
        'G': 18,  # تاريخ التقديم
        'H': 18,  # تاريخ المراجعة
        'I': 20,  # راجعه
        'J': 15,  # الهاتف
        'K': 25,  # البريد
        'L': 25,  # الكلية
        'M': 25,  # القسم
        'N': 15,  # السنة
    }
    
    for col, width in column_widths.items():
        ws.column_dimensions[col].width = width
    
    # ارتفاع الصف الأول (العناوين)
    ws.row_dimensions[1].height = 25
    
    # ==========================================
    # 7. إضافة معلومات إضافية (اختياري)
    # ==========================================
    
    # صف فارغ
    summary_row = applications.count() + 3
    
    ws.cell(row=summary_row, column=1).value = "الإجمالي:"
    ws.cell(row=summary_row, column=2).value = applications.count()
    ws.cell(row=summary_row, column=1).font = Font(bold=True)
    ws.cell(row=summary_row, column=2).font = Font(bold=True)
    
    ws.cell(row=summary_row + 1, column=1).value = "تاريخ التصدير:"
    ws.cell(row=summary_row + 1, column=2).value = timezone.now().strftime('%Y-%m-%d %H:%M')
    
    ws.cell(row=summary_row + 2, column=1).value = "صُدّر بواسطة:"
    ws.cell(row=summary_row + 2, column=2).value = request.user.get_full_name() or request.user.username
    
    # ==========================================
    # 8. حفظ الملف وإرجاعه
    # ==========================================
    
    # اسم الملف
    filename = f"applications_{timezone.now().strftime('%Y-%m-%d_%H-%M')}.xlsx"
    
    # إنشاء HttpResponse مع نوع المحتوى
    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    
    # حفظ الـ Workbook في الـ response
    wb.save(response)
    
    return response
    # ========================================
# Views الإشعارات المتقدمة
# ========================================

@never_cache
@login_required(login_url='students:login')
def notifications_list(request):
    """صفحة كل الإشعارات"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # جلب كل الإشعارات
    notifications = student.student_notifications.all()
    
    # الفلترة حسب الحالة
    filter_type = request.GET.get('filter', 'all')
    
    if filter_type == 'unread':
        notifications = notifications.filter(is_read=False)
    elif filter_type == 'read':
        notifications = notifications.filter(is_read=True)
    
    # إحصائيات
    total_count = student.student_notifications.count()
    unread_count = student.student_notifications.filter(is_read=False).count()
    read_count = student.student_notifications.filter(is_read=True).count()
    
    context = {
        'student': student,
        'notifications': notifications,
        'total_count': total_count,
        'unread_count': unread_count,
        'read_count': read_count,
        'filter_type': filter_type,
        'page_title': 'إشعاراتي'
    }
    
    return render(request, 'students/notifications_list.html', context)


@never_cache
@login_required(login_url='students:login')
@require_http_methods(["POST"])
def mark_notification_read(request, pk):
    """تمييز إشعار كمقروء"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'لا يوجد ملف طالب'})
    
    try:
        notification = get_object_or_404(
            StudentNotification,
            id=pk,
            student=student
        )
        
        notification.mark_as_read()
        
        # إرجاع عدد الإشعارات غير المقروءة المحدث
        unread_count = student.student_notifications.filter(is_read=False).count()
        
        return JsonResponse({
            'success': True,
            'unread_count': unread_count
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@never_cache
@login_required(login_url='students:login')
@require_http_methods(["POST"])
def mark_all_notifications_read(request):
    """تمييز كل الإشعارات كمقروءة"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'لا يوجد ملف طالب'})
    
    try:
        # تحديث كل الإشعارات غير المقروءة
        unread_notifications = student.student_notifications.filter(is_read=False)
        count = unread_notifications.count()
        
        for notification in unread_notifications:
            notification.mark_as_read()
        
        return JsonResponse({
            'success': True,
            'count': count,
            'message': f'تم تمييز {count} إشعار كمقروء'
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@never_cache
@login_required(login_url='students:login')
@require_http_methods(["POST"])
def delete_notification(request, pk):
    """حذف إشعار"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'لا يوجد ملف طالب'})
    
    try:
        notification = get_object_or_404(
            StudentNotification,
            id=pk,
            student=student
        )
        
        notification.delete()
        
        # إرجاع عدد الإشعارات غير المقروءة المحدث
        unread_count = student.student_notifications.filter(is_read=False).count()
        
        return JsonResponse({
            'success': True,
            'unread_count': unread_count,
            'message': 'تم حذف الإشعار'
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})


@never_cache
@login_required(login_url='students:login')
@require_http_methods(["POST"])
def delete_read_notifications(request):
    """حذف كل الإشعارات المقروءة"""
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'لا يوجد ملف طالب'})
    
    try:
        read_notifications = student.student_notifications.filter(is_read=True)
        count = read_notifications.count()
        read_notifications.delete()
        
        return JsonResponse({
            'success': True,
            'count': count,
            'message': f'تم حذف {count} إشعار مقروء'
        })
        
    except Exception as e:
        return JsonResponse({'success': False, 'error': str(e)})
    # ========================================
# API للإشعارات Real-Time
# ========================================

@never_cache
@login_required(login_url='students:login')
def notifications_api(request):
    """
    API لجلب الإشعارات - للتحديث التلقائي
    
    يُستدعى كل 30 ثانية من JavaScript
    يُرجع: عدد الإشعارات + قائمة آخر 5
    """
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        return JsonResponse({'success': False, 'error': 'لا يوجد ملف طالب'})
    
    # عدد الإشعارات غير المقروءة
    unread_count = student.get_unread_notifications_count()
    
    # آخر 5 إشعارات
    recent_notifications = student.get_recent_notifications()
    
    # تحويل الإشعارات إلى JSON
    notifications_data = []
    for notif in recent_notifications:
        notifications_data.append({
            'id': notif.id,
            'title': notif.title,
            'message': notif.message,
            'link': notif.link or '',
            'is_read': notif.is_read,
            'icon': notif.get_icon(),
            'color': notif.get_color(),
            'created_at': notif.created_at.isoformat(),
            'time_ago': f'{notif.created_at}',  # سيُحسب في JS
        })
    
    return JsonResponse({
        'success': True,
        'unread_count': unread_count,
        'total_count': student.student_notifications.count(),
        'notifications': notifications_data
    })
# ========================================
# Views إدارة الحساب الشخصي
# ========================================

@never_cache
@login_required(login_url='students:login')
def profile_view(request):
    """
    صفحة عرض المعلومات الشخصية
    
    ماذا تفعل؟
    - تعرض كل معلومات الطالب
    - تصميم بطاقة احترافية
    """
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # إحصائيات الطالب
    total_applications = student.applications.count()
    approved_count = student.applications.filter(status='approved').count()
    pending_count = student.applications.filter(status='pending').count()
    
    context = {
        'student': student,
        'total_applications': total_applications,
        'approved_count': approved_count,
        'pending_count': pending_count,
        'page_title': 'معلوماتي الشخصية'
    }
    
    return render(request, 'students/profile/view.html', context)


@never_cache
@login_required(login_url='students:login')
def profile_edit(request):
    """
    صفحة تعديل المعلومات الشخصية
    """
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    if request.method == 'POST':
        # نموذجان: واحد للـ Profile، واحد للـ User
        profile_form = StudentProfileUpdateForm(
            request.POST, 
            request.FILES,  # ← مهم للصور!
            instance=student
        )
        email_form = UserEmailUpdateForm(request.POST, instance=request.user)
        
        if profile_form.is_valid() and email_form.is_valid():
            try:
                # حفظ التعديلات
                profile_form.save()
                email_form.save()
                
                # إنشاء إشعار
                StudentNotification.objects.create(
                    student=student,
                    title='تم تحديث معلوماتك ✅',
                    message='تم تحديث معلوماتك الشخصية بنجاح.',
                    notification_type='success',
                    link='/students/profile/'
                )
                
                messages.success(request, '✅ تم تحديث معلوماتك بنجاح!')
                return redirect('students:profile_view')
                    
            except Exception as e:
                messages.error(request, f'❌ حدث خطأ: {str(e)}')
        else:
            # عرض الأخطاء
            for field, errors in profile_form.errors.items():
                for error in errors:
                    messages.error(request, f'❌ {error}')
            for field, errors in email_form.errors.items():
                for error in errors:
                    messages.error(request, f'❌ {error}')
    else:
        # GET: عرض النموذج بالمعلومات الحالية
        profile_form = StudentProfileUpdateForm(instance=student)
        email_form = UserEmailUpdateForm(instance=request.user)
    
    context = {
        'student': student,
        'profile_form': profile_form,
        'email_form': email_form,
        'page_title': 'تعديل معلوماتي'
    }
    
    return render(request, 'students/profile/edit.html', context)


@never_cache
@login_required(login_url='students:login')
def change_password(request):
    """
    صفحة تغيير كلمة المرور
    """
    
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    if request.method == 'POST':
        form = PasswordChangeForm(request.POST)
        
        if form.is_valid():
            old_password = form.cleaned_data['old_password']
            new_password = form.cleaned_data['new_password1']
            
            # التحقق من كلمة المرور القديمة
            if not request.user.check_password(old_password):
                messages.error(request, '❌ كلمة المرور القديمة غير صحيحة')
            else:
                try:
                    # حفظ كلمة المرور الجديدة
                    request.user.set_password(new_password)
                    request.user.save()
                    
                    # مهم جداً! للحفاظ على الجلسة
                    update_session_auth_hash(request, request.user)
                    
                    # إنشاء إشعار
                    StudentNotification.objects.create(
                        student=student,
                        title='تم تغيير كلمة المرور 🔒',
                        message='تم تغيير كلمة المرور الخاصة بك بنجاح.',
                        notification_type='success'
                    )
                    
                    messages.success(request, '✅ تم تغيير كلمة المرور بنجاح!')
                    return redirect('students:profile_view')
                    
                except Exception as e:
                    messages.error(request, f'❌ حدث خطأ: {str(e)}')
        else:
            # عرض أخطاء النموذج
            for field, errors in form.errors.items():
                for error in errors:
                    messages.error(request, f'❌ {error}')
    else:
        form = PasswordChangeForm()
    
    context = {
        'student': student,
        'form': form,
        'page_title': 'تغيير كلمة المرور'
    }
    
    return render(request, 'students/profile/change_password.html', context)


@never_cache
@login_required(login_url='students:login')
@require_http_methods(["POST"])
def delete_profile_image(request):
    """
    حذف الصورة الشخصية
    
    AJAX endpoint
    """
    try:
        student = request.user.student_profile
        
        if student.photo:  # ← غيرنا من profile_image إلى photo
            # حذف الملف من السيرفر
            student.photo.delete()  # ← وهنا أيضاً
            student.save()
            
            return JsonResponse({
                'success': True,
                'message': 'تم حذف الصورة بنجاح'
            })
        else:
            return JsonResponse({
                'success': False,
                'error': 'لا توجد صورة لحذفها'
            })
            
    except Exception as e:
        return JsonResponse({
            'success': False,
            'error': str(e)
        })
    

    # ========================================
# Views الطباعة والتصدير
# ========================================

@never_cache
@login_required(login_url='students:login')
def print_application(request, application_id):
    """
    صفحة طباعة طلب واحد
    """
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # التحقق من ملكية الطلب
    application = get_object_or_404(
        StudentApplication,
        id=application_id,
        student=student
    )
    
    context = {
        'application': application,
        'student': student,
        'page_title': f'طباعة الطلب {application.application_number}'
    }
    
    return render(request, 'students/applications/print_application.html', context)

@never_cache
@login_required(login_url='students:login')
def print_all_applications(request):
    """
    صفحة طباعة كل الطلبات
    """
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    applications = student.applications.all().order_by('-submitted_at')
    
    context = {
        'applications': applications,
        'student': student,
        'page_title': 'طباعة كل الطلبات'
    }
    
    return render(request, 'students/print_all.html', context)  # ← المسار الصحيح



@never_cache
@login_required(login_url='students:login')
def student_card(request):
    """
    بطاقة الطالب الرقمية
    """
    try:
        student = request.user.student_profile
    except StudentProfile.DoesNotExist:
        messages.error(request, 'لا يوجد ملف طالب لهذا المستخدم')
        return redirect('homepage:index')
    
    # توليد QR Code (البيانات)
    qr_data = f"Student:{student.registration_number}|Name:{student.full_name_arabic}|Faculty:{student.faculty}"
    
    context = {
        'student': student,
        'qr_data': qr_data,
        'page_title': 'بطاقة الطالب الرقمية'
    }
    
    return render(request, 'students/profile/student_card.html', context)


@never_cache
@login_required(login_url='students:login')
def download_application_pdf(request, application_id):
    """
    تحميل طلب كـ PDF
    
    ملاحظة: يتطلب تثبيت مكتبة xhtml2pdf أو WeasyPrint
    للبساطة الآن: سنعيد توجيه لصفحة الطباعة
    يمكن للطالب استخدام "طباعة إلى PDF" من المتصفح
    """
    # إعادة توجيه لصفحة الطباعة مع رسالة
    messages.info(request, 'استخدم "طباعة" من المتصفح واختر "حفظ كـ PDF"')
    return redirect('students:print_application', application_id=application_id)


@never_cache
@login_required(login_url='students:login')
def export_all_pdf(request):
    """
    تصدير كل الطلبات كـ PDF
    """
    messages.info(request, 'استخدم "طباعة" من المتصفح واختر "حفظ كـ PDF"')
    return redirect('students:print_all_applications')